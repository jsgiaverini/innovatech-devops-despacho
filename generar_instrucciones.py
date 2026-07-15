from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

def add_heading_custom(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    return h

def add_para(text, bold=False, italic=False, size=11, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_bullet(text, bold=False):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()
    return table

# ═══════════ PORTADA ═══════════
for _ in range(8):
    doc.add_paragraph()

add_para('INSTRUCCIONES PARA CAPTURA DE EVIDENCIAS', bold=True, size=22,
         color=RGBColor(0xC0, 0x39, 0x2B), align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_para('Proyecto Innovatech Chile - Sistema de Gestión de Despachos', bold=True, size=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
add_para('Evaluación Parcial N°2 - Introducción a Herramientas DevOps (ISY1101)', size=12,
         align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_para('Junio 2026', size=12, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# ═══════════ INTRODUCCIÓN ═══════════
add_heading_custom('Introducción', level=1)
add_para(
    'Este documento contiene las instrucciones paso a paso para que puedas tomar las 12 capturas de '
    'pantalla necesarias como evidencia del despliegue completo del proyecto Innovatech en AWS EC2.'
)
add_para(
    'Una vez tomadas las capturas, debes insertarlas en el informe principal '
    '"INFORME_EVIDENCIAS_Innovatech_DevOps.docx" en las secciones marcadas con [INSERTAR CAPTURA DE PANTALLA AQUÍ].')
doc.add_paragraph()

add_para(
    'IMPORTANTE: Las capturas que requieren consola AWS (EC2, ECR) debes tomarlas desde el '
    'AWS Academy Learner Lab. Las credenciales AWS son temporales y rotan cada vez que inicias el laboratorio.',
    bold=True, size=11, color=RGBColor(0xC0, 0x39, 0x2B)
)

# ═══════════ LISTA DE CAPTURAS ═══════════
doc.add_page_break()
add_heading_custom('Lista de Capturas a Realizar', level=1)

# ─── 1 ───
add_heading_custom('Captura 1: GitHub Actions - 3 Workflows en Verde', level=2)
add_para(
    'Accede a: https://github.com/jsgiaverini/innovatech-devops-despacho/actions',
    italic=True
)
add_para('Pasos:', bold=True)
add_bullet('Inicia sesión en GitHub')
add_bullet('Ve al repositorio: jsgiaverini/innovatech-devops-despacho')
add_bullet('Haz clic en la pestaña "Actions"')
add_bullet('Verifica que los 3 workflows tengan check verde (success):')
add_bullet('  • Frontend CI/CD - Build, Push & Deploy')
add_bullet('  • Backend Ventas CI/CD - Build, Push & Deploy')
add_bullet('  • Backend Despachos CI/CD - Build, Push & Deploy')
add_bullet('Toma la captura de la pantalla completa mostrando la lista de runs con los 3 checks verdes')
add_bullet('Nombra el archivo como: 01-github-actions.png')
doc.add_paragraph()

# ─── 2 ───
add_heading_custom('Captura 2: Frontend Cargando en Navegador', level=2)
add_para('Accede a: http://3.227.232.165', italic=True)
add_para('Pasos:', bold=True)
add_bullet('Abre el navegador (Chrome, Firefox, Edge)')
add_bullet('Ingresa la URL: http://3.227.232.165')
add_bullet('Espera a que la aplicación cargue completamente')
add_bullet('Toma la captura mostrando la página principal de Innovatech funcionando')
add_bullet('Nombra el archivo como: 02-frontend.png')
doc.add_paragraph()

# ─── 3 ───
add_heading_custom('Captura 3: API Ventas - Respuesta Correcta', level=2)
add_para('Accede a: http://3.227.232.165/api/v1/ventas', italic=True)
add_para('Pasos:', bold=True)
add_bullet('En la misma pestaña del navegador o una nueva')
add_bullet('Ingresa: http://3.227.232.165/api/v1/ventas')
add_bullet('Debe mostrar: [] (arreglo JSON vacío)')
add_bullet('Toma la captura mostrando la respuesta JSON en el navegador')
add_bullet('Nombra el archivo como: 03-api-ventas.png')
doc.add_paragraph()

# ─── 4 ───
add_heading_custom('Captura 4: API Despachos - Respuesta Correcta', level=2)
add_para('Accede a: http://3.227.232.165/api/v1/despachos', italic=True)
add_para('Pasos:', bold=True)
add_bullet('Ingresa: http://3.227.232.165/api/v1/despachos')
add_bullet('Debe mostrar: [] (arreglo JSON vacío)')
add_bullet('Toma la captura mostrando la respuesta JSON en el navegador')
add_bullet('Nombra el archivo como: 04-api-despachos.png')
doc.add_paragraph()

# ─── 5 ───
add_heading_custom('Captura 5: EC2 - Instancias Ejecutándose', level=2)
add_para('Accede a: Consola AWS > EC2 > Instances', italic=True)
add_para('Pasos:', bold=True)
add_bullet('Inicia sesión en AWS Academy Learner Lab')
add_bullet('Abre la consola AWS y ve a EC2')
add_bullet('Selecciona "Instances" en el menú lateral')
add_bullet('Verifica que aparezcan 2 instancias:')
add_bullet('  • innovatech-frontend (Estado: Running)')
add_bullet('  • innovatech-backend (Estado: Running)')
add_bullet('Toma la captura mostrando la lista de instancias con sus estados, IPs y tipos')
add_bullet('Nombra el archivo como: 05-ec2-instancias.png')
doc.add_paragraph()

# ─── 6 ───
add_heading_custom('Captura 6: EC2 - Contenedores en Ejecución (SSH)', level=2)
add_para('Comando SSH en ambas instancias', italic=True)
add_para('Pasos:', bold=True)
add_bullet('Abre la terminal (PowerShell, CMD o Git Bash)')
add_para('Instancia Frontend:', bold=True)
add_para('ssh -i "tu-llave.pem" ec2-user@<IP-FRONTEND>', size=9, italic=True)
add_bullet('Ejecuta: sudo docker ps')
add_bullet('Debe mostrar el contenedor "front_despacho" corriendo')
add_bullet('Toma captura de la salida')
doc.add_paragraph()
add_para('Instancia Backend:', bold=True)
add_para('ssh -i "tu-llave.pem" ec2-user@<IP-BACKEND>', size=9, italic=True)
add_bullet('Ejecuta: sudo docker ps')
add_bullet('Debe mostrar 3 contenedores: innovatech_mysql, back_ventas, back_despachos')
add_bullet('Toma captura de la salida')
add_bullet('Nombra los archivos como: 06a-docker-ps-frontend.png y 06b-docker-ps-backend.png')
doc.add_paragraph()

# ─── 7 ───
add_heading_custom('Captura 7: ECR - Repositorios con Imágenes', level=2)
add_para('Accede a: Consola AWS > ECR > Repositories', italic=True)
add_para('Pasos:', bold=True)
add_bullet('En la consola AWS, busca y abre "Elastic Container Registry (ECR)"')
add_bullet('Ve a "Repositories" en el menú lateral')
add_bullet('Verifica que aparezcan 3 repositorios:')
add_bullet('  • front-despacho (con al menos 1 imagen)')
add_bullet('  • back-ventas (con al menos 1 imagen)')
add_bullet('  • back-despachos (con al menos 1 imagen)')
add_bullet('Toma la captura de la lista de repositorios ECR')
add_bullet('Opcional: abre un repositorio para mostrar la imagen "latest"')
add_bullet('Nombra el archivo como: 07-ecr-repositorios.png')
doc.add_paragraph()

# ─── 8 ───
add_heading_custom('Captura 8: GitHub Secrets Configurados', level=2)
add_para('Accede a: GitHub > Settings > Secrets and variables > Actions', italic=True)
add_para('Pasos:', bold=True)
add_bullet('En GitHub, ve al repositorio jsgiaverini/innovatech-devops-despacho')
add_bullet('Haz clic en "Settings"')
add_bullet('Menú lateral: "Secrets and variables" > "Actions"')
add_bullet('Verifica que aparezcan los secrets (NO es necesario mostrar los valores)')
add_bullet('Debe mostrar al menos: AWS_ACCESS_KEY_ID, AWS_SECRET_ACCESS_KEY, AWS_REGION,')
add_bullet('  ECR_REGISTRY, EC2_FRONTEND_HOST, EC2_BACKEND_HOST, EC2_SSH_KEY, etc.')
add_bullet('Toma la captura de la pantalla de Secrets')
add_bullet('Nombra el archivo como: 08-github-secrets.png')
doc.add_paragraph()

# ─── 9 ───
add_heading_custom('Captura 9: Dockerfiles Multi-Stage', level=2)
add_para('Código fuente en el repositorio', italic=True)
add_para('Pasos:', bold=True)
add_bullet('En el repositorio de GitHub, navega a:')
add_bullet('  • front_despacho/Dockerfile')
add_bullet('  • back-Ventas_SpringBoot/Springboot-API-REST/Dockerfile')
add_bullet('  • back-Despachos_SpringBoot/Springboot-API-REST-DESPACHO/Dockerfile')
add_bullet('Toma captura del contenido de cada Dockerfile (o una captura combinada)')
add_bullet('Debe verse: multi-stage (FROM ... AS build), usuario no root, HEALTHCHECK')
add_bullet('Nombra los archivos como: 09a-dockerfile-frontend.png, 09b-dockerfile-ventas.png, 09c-dockerfile-despachos.png')
doc.add_paragraph()

# ─── 10 ───
add_heading_custom('Captura 10: docker-compose.yml - Stack Completo', level=2)
add_para('Archivo raíz del proyecto', italic=True)
add_para('Pasos:', bold=True)
add_bullet('En el repositorio, abre el archivo docker-compose.yml (raíz)')
add_bullet('Toma captura mostrando:')
add_bullet('  • Los 4 servicios: mysql, back-ventas, back-despachos, frontend')
add_bullet('  • La red interna: innovatech_net')
add_bullet('  • El volumen nombrado: mysql_data')
add_bullet('  • Las dependencias con healthcheck')
add_bullet('Nombra el archivo como: 10-docker-compose.png')
doc.add_paragraph()

# ─── 11 ───
add_heading_custom('Captura 11: README.md - Documentación', level=2)
add_para('Archivo README.md del proyecto', italic=True)
add_para('Pasos:', bold=True)
add_bullet('En el repositorio, abre el archivo README.md (raíz)')
add_bullet('Toma 2-3 capturas que muestren:')
add_bullet('  • Diagrama de arquitectura')
add_bullet('  • Estructura del proyecto')
add_bullet('  • Instrucciones de despliegue AWS')
add_bullet('Nombra los archivos como: 11a-readme-1.png, 11b-readme-2.png, etc.')
doc.add_paragraph()

# ─── 12 ───
add_heading_custom('Captura 12: Git Log - Commits Descriptivos', level=2)
add_para('Terminal local o GitHub Insights', italic=True)
add_para('Pasos:', bold=True)
add_bullet('Opción 1 - Terminal local:')
add_bullet('  Abre la terminal en la carpeta del proyecto')
add_bullet('  Ejecuta: git log --oneline -10')
add_bullet('  Toma captura de la salida mostrando los 7 commits')
add_bullet('Opción 2 - GitHub:')
add_bullet('  En el repositorio, ve a "Insights" > "Network"')
add_bullet('  O ve directamente al historial de commits')
add_bullet('Nombra el archivo como: 12-git-log.png')
doc.add_paragraph()

# ═══════════ INSTRUCCIONES FINALES ═══════════
doc.add_page_break()
add_heading_custom('Instrucciones Finales', level=1)

add_para('Después de tomar todas las capturas:', bold=True, size=12)
add_para('')
add_bullet('Abre el archivo: INFORME_EVIDENCIAS_Innovatech_DevOps.docx (está en la carpeta Evidencias/)')
add_bullet('Busca cada sección que dice "[INSERTAR CAPTURA DE PANTALLA AQUÍ]"')
add_bullet('Inserta la imagen correspondiente usando: Insertar > Imágenes (Word)')
add_bullet('Ajusta el tamaño de la imagen para que se vea completa en la página')
add_bullet('Guarda el documento')

add_para('')
add_para('Resumen de archivos a generar:', bold=True, size=12)

add_table(
    ['#', 'Nombre Archivo', 'Descripción'],
    [
        ['1', '01-github-actions.png', 'GitHub Actions - 3 workflows verdes'],
        ['2', '02-frontend.png', 'Frontend en http://3.227.232.165'],
        ['3', '03-api-ventas.png', 'API Ventas respondiendo []'],
        ['4', '04-api-despachos.png', 'API Despachos respondiendo []'],
        ['5', '05-ec2-instancias.png', 'EC2 instancias Running'],
        ['6a', '06a-docker-ps-frontend.png', 'docker ps en EC2 Frontend'],
        ['6b', '06b-docker-ps-backend.png', 'docker ps en EC2 Backend'],
        ['7', '07-ecr-repositorios.png', 'ECR - 3 repositorios'],
        ['8', '08-github-secrets.png', 'GitHub Secrets'],
        ['9a', '09a-dockerfile-frontend.png', 'Dockerfile Frontend'],
        ['9b', '09b-dockerfile-ventas.png', 'Dockerfile Backend Ventas'],
        ['9c', '09c-dockerfile-despachos.png', 'Dockerfile Backend Despachos'],
        ['10', '10-docker-compose.png', 'docker-compose.yml raíz'],
        ['11a', '11a-readme-1.png', 'README.md (parte 1)'],
        ['11b', '11b-readme-2.png', 'README.md (parte 2)'],
        ['12', '12-git-log.png', 'Git log - commits'],
    ]
)

# ─── Guardar ───
output_path = os.path.join(os.path.dirname(__file__), 'INSTRUCCIONES_CAPTURA_EVIDENCIAS.docx')
doc.save(output_path)
print(f'Documento generado: {output_path}')
